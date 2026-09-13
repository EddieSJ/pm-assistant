"""Word 文档生成：把 Markdown 文本渲染为 .docx 并可下载。"""
from pathlib import Path

from docx import Document

from config import Config


def markdown_to_word(title: str, markdown_text: str, filename: str | None = None,
                     version: str | None = None) -> Path:
    doc = Document()
    doc.add_heading(title, level=0)

    for block in markdown_text.split("\n"):
        block = block.rstrip()
        if not block.strip():
            continue
        if block.startswith("### "):
            doc.add_heading(block[4:].strip(), level=3)
        elif block.startswith("## "):
            doc.add_heading(block[3:].strip(), level=2)
        elif block.startswith("# "):
            doc.add_heading(block[2:].strip(), level=1)
        elif block.startswith("- ") or block.startswith("* "):
            doc.add_paragraph(block[2:].strip(), style="List Bullet")
        elif block.lstrip().startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")):
            doc.add_paragraph(block.lstrip(), style="List Number")
        elif block.startswith("|"):
            # 简单表格行按文本保留
            doc.add_paragraph(block)
        else:
            doc.add_paragraph(block)

    if version:
        safe_title = title.strip().replace(" ", "_")
        filename = f"{safe_title}_{version}.docx"
    else:
        filename = (filename or f"{title}.docx").strip()
        if not filename.lower().endswith(".docx"):
            filename += ".docx"
    out_dir = Config.output_dir / "documents"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / filename
    doc.save(str(path))
    return path


def build_summary_markdown(project: dict, facts: list[dict], stakeholders: list[dict],
                           finance: list[dict], actions: list[dict], conflicts: list[dict]) -> str:
    """基于看板数据生成一份项目综合整理报告的 Markdown。"""
    lines = [
        f"# {project.get('name', '项目')} 综合整理报告",
        "",
        f"项目类型：{project.get('project_type') or '未识别'}",
        f"项目编码：{project.get('project_code') or '-'}",
        f"生成时间：{__import__('db').now()}",
        "",
        "## 一、项目事实台账",
    ]
    for f in facts:
        lines.append(f"- 【{f['category']}】{f['field']}：{f['value']}（{f['status']}）[来源:{f['source']}]")
    lines += ["", "## 二、人员与干系人"]
    for s in stakeholders:
        lines.append(f"- {s['name']}（{s['organization']} / {s['role']}）")
    lines += ["", "## 三、财务信息"]
    for f in finance:
        lines.append(f"- 【{f['category']}】{f['item']}：{f['amount']}（{f['tax_inclusive']} / {f['period']}）[来源:{f['source']}]")
    lines += ["", "## 四、行动项"]
    for a in actions:
        lines.append(f"- {a['action']}（负责人:{a['owner']}，期限:{a['due_date']}，状态:{a['status']}）")
    lines += ["", "## 五、冲突与待确认事项"]
    for c in conflicts:
        lines.append(f"- 【{c['category']}】{c['description']}")
    lines += ["", "> 说明：本报告由私人智能体整理，财务/合同/人事内容需人工复核。"]
    return "\n".join(lines)

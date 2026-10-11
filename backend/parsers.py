"""文件解析器：把 .md/.pdf/.docx/.xlsx 统一抽取为纯文本。"""
from pathlib import Path


def parse_md(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def parse_pdf(path: Path) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(str(path))
    try:
        parts = []
        for page in doc:
            parts.append(page.get_text())
        return "\n".join(parts)
    finally:
        doc.close()


def parse_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        parts.append("")  # 空行分隔表格
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"【工作表：{ws.title}】")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


def parse_file(path: str | Path) -> tuple[str, str]:
    """解析单个文件，返回 (文本内容, 扩展名)。"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".md":
        return parse_md(p), "md"
    if ext == ".pdf":
        return parse_pdf(p), "pdf"
    if ext == ".docx":
        return parse_docx(p), "docx"
    if ext == ".xlsx":
        return parse_xlsx(p), "xlsx"
    # 其他文本类直接读
    try:
        return p.read_text(encoding="utf-8", errors="replace"), ext.lstrip(".")
    except Exception:
        return "", ext.lstrip(".")


SUPPORTED_EXTS = {".md", ".pdf", ".docx", ".xlsx", ".txt", ".csv"}


def parse_csv(path: Path) -> str:
    """CSV 转纯文本（utf-8 优先，gbk 兜底），沿用文本管道。"""
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
        except Exception:
            return ""
    return ""


def parse_table(path: str | Path) -> dict:
    """结构化解析表格文件，返回 {"columns": [...], "rows": [[...]], "sheet": str}。

    支持 .csv（utf-8 + gbk 兜底）与 .xlsx（openpyxl，取第一个 sheet）。
    空表/无表头返回 {"columns": [], "rows": [], "error": "..."}。
    """
    import csv as _csv
    import io

    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        text = ""
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                text = p.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
            except Exception:
                text = ""
                break
        if not text.strip():
            return {"columns": [], "rows": [], "error": "空文件"}
        reader = _csv.reader(io.StringIO(text))
        matrix = [[(c or "").strip() for c in row] for row in reader if any((c or "").strip() for c in row)]
        if not matrix:
            return {"columns": [], "rows": [], "error": "空表"}
        return {"columns": matrix[0], "rows": matrix[1:], "sheet": p.name}
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(p), data_only=True, read_only=True)
            try:
                if not wb.worksheets:
                    return {"columns": [], "rows": [], "error": "无工作表"}
                ws = wb.worksheets[0]
                sheet_name = ws.title  # close 前缓存，避免 close 后引用 ws
                matrix = []
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c).strip() for c in row]
                    if any(cells):
                        matrix.append(cells)
            finally:
                wb.close()
        except Exception as e:
            return {"columns": [], "rows": [], "error": f"解析失败: {e}"}
        if not matrix:
            return {"columns": [], "rows": [], "error": "空表"}
        return {"columns": matrix[0], "rows": matrix[1:], "sheet": sheet_name if wb else p.name}
    return {"columns": [], "rows": [], "error": f"不支持的类型: {ext}"}
